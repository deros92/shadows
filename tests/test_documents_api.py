import io
import json

import docx
import pytest
from fastapi.testclient import TestClient

from shadows.audit import AuditLog
from shadows.documents import UnsupportedDocument, extract_text
from shadows.server import create_app


def make_pdf(lines: list[str]) -> bytes:
    """PDF minimo con testo vero (Helvetica), senza dipendenze di scrittura PDF."""
    ops = ["BT", "/F1 12 Tf", "14 TL", "72 760 Td"]
    for line in lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops.append(f"({safe}) Tj T*")
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def make_docx(paragraph: str, cell: str) -> bytes:
    document = docx.Document()
    document.add_paragraph(paragraph)
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = cell
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_extract_pdf_docx_txt():
    pdf = extract_text("contratto.pdf", make_pdf(["Cliente SuperMercati Bianchi", "IBAN IT60X0542811101000000123456"]))
    assert pdf.kind == "pdf" and pdf.pages == 1
    assert "SuperMercati Bianchi" in pdf.text and "IT60X0542811101000000123456" in pdf.text

    word = extract_text("policy.docx", make_docx("Referente a@b.it", "Rossi Food Service"))
    assert "a@b.it" in word.text and "Rossi Food Service" in word.text

    assert extract_text("note.txt", "perché".encode("cp1252")).text == "perché"


def test_unsupported_format():
    with pytest.raises(UnsupportedDocument):
        extract_text("foglio.xlsx", b"PK")


@pytest.fixture
def client(engine, tmp_path):
    audit_path = tmp_path / "audit.jsonl"
    app = create_app(engine, AuditLog(audit_path))
    c = TestClient(app, base_url="http://127.0.0.1:8765")
    c.audit_path = audit_path
    return c


def test_analyze_restore_and_audit_has_no_text(client):
    text = "Contratto con SuperMercati Bianchi, mail a@b.it"
    resp = client.post("/v1/analyze", json={"text": text, "site": "chatgpt.com"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["action"] == "redact"
    assert body["anonymized_text"] == "Contratto con [CLIENTE_1], mail [EMAIL_1]"
    assert {e["label"]: e["optional"] for e in body["entities"]} == {"CLIENTE": False, "EMAIL": False}

    restored = client.post("/v1/restore", json={"text": "Per [CLIENTE_1] scrivo a [EMAIL_1]",
                                                "mapping": body["mapping"]}).json()["text"]
    assert restored == "Per SuperMercati Bianchi scrivo a a@b.it"

    log = client.audit_path.read_text(encoding="utf-8")
    record = json.loads(log.splitlines()[0])
    assert record["labels"] == {"CLIENTE": 1, "EMAIL": 1}
    assert "Bianchi" not in log and "a@b.it" not in log


def test_entity_offsets_are_code_points(client):
    text = "🙂 mail a@b.it"
    ent = client.post("/v1/analyze", json={"text": text}).json()["entities"][0]
    assert "".join(list(text)[ent["start"]:ent["end"]]) == "a@b.it"


def test_document_endpoint(client):
    pdf = make_pdf(["Cliente SuperMercati Bianchi, CF RSSMRA85M01H501Q"])
    resp = client.post("/v1/document", files={"file": ("c.pdf", pdf, "application/pdf")},
                       data={"known_mapping": json.dumps({"[CLIENTE_1]": "Rossi Food Service"})})
    assert resp.status_code == 200
    body = resp.json()
    assert body["document"]["kind"] == "pdf"
    assert "[CLIENTE_2]" in body["anonymized_text"] and "[CF_1]" in body["anonymized_text"]


def test_rejects_web_origins_and_foreign_hosts(client):
    ok = client.post("/v1/analyze", json={"text": "ciao"},
                     headers={"Origin": "chrome-extension://abcdef"})
    assert ok.status_code == 200
    evil = client.post("/v1/analyze", json={"text": "ciao"}, headers={"Origin": "https://evil.example"})
    assert evil.status_code == 403
    rebinding = client.get("/health", headers={"Host": "attacker.example:8765"})
    assert rebinding.status_code == 403


def test_demo_page_served(client):
    resp = client.get("/demo")
    assert resp.status_code == 200 and "Chat di prova" in resp.text
