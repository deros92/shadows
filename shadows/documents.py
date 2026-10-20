"""Estrazione del testo dai documenti con sole librerie a licenza permissiva.

All'AI arriva il testo anonimizzato, non il file: basta quindi leggere i documenti.
PDF con pdfplumber (MIT), Word con python-docx (MIT), OCR facoltativo con
Tesseract (Apache-2.0) sulle pagine renderizzate da pypdfium2 (Apache-2.0/BSD).
"""

from __future__ import annotations

import io
import os
import shutil
from dataclasses import dataclass, field
from pathlib import Path, PurePath

MAX_BYTES = 20 * 1024 * 1024
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".log", ".eml"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
MIN_PAGE_CHARS = 25  # sotto questa soglia una pagina PDF è considerata una scansione

# Modelli linguistici nel progetto (vedi README), così non servono permessi di
# amministratore per aggiungere l'italiano all'installazione di Tesseract.
TESSDATA_DIR = Path(__file__).resolve().parent.parent / "tessdata"
_TESSERACT_CANDIDATES = [
    Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "Tesseract-OCR" / "tesseract.exe",
    Path(os.environ.get("LOCALAPPDATA", "")) / "Programs" / "Tesseract-OCR" / "tesseract.exe",
]


class UnsupportedDocument(ValueError):
    pass


@dataclass
class ExtractedDocument:
    text: str
    kind: str
    pages: int = 0
    warnings: list[str] = field(default_factory=list)


def extract_text(filename: str, data: bytes) -> ExtractedDocument:
    if len(data) > MAX_BYTES:
        raise UnsupportedDocument(f"file oltre {MAX_BYTES // (1024 * 1024)} MB")
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".pdf":
        return _pdf(data)
    if suffix == ".docx":
        return _docx(data)
    if suffix in TEXT_SUFFIXES:
        return ExtractedDocument(_decode(data), kind="testo")
    if suffix in IMAGE_SUFFIXES:
        return _image(data)
    raise UnsupportedDocument(f"formato non supportato: {suffix or 'senza estensione'}")


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _pdf(data: bytes) -> ExtractedDocument:
    import pdfplumber

    texts: list[str] = []
    empty_pages: list[int] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        n_pages = len(pdf.pages)
        for i, page in enumerate(pdf.pages, start=1):
            page_text = page.extract_text() or ""
            if len(page_text.strip()) < MIN_PAGE_CHARS:
                empty_pages.append(i)
            texts.append(page_text)

    doc = ExtractedDocument("", kind="pdf", pages=n_pages)
    if empty_pages:
        if ocr_available():
            for i, ocr_text in _ocr_pages(data, empty_pages):
                # Si tiene il testo più lungo: una pagina con solo un'intestazione
                # digitale e il corpo scansionato non perde nessuna delle due parti.
                if len(ocr_text.strip()) > len(texts[i - 1].strip()):
                    texts[i - 1] = ocr_text
            doc.warnings.append(f"OCR applicato alle pagine {empty_pages}: verificare il testo, "
                                "l'OCR può sbagliare qualche carattere")
        else:
            doc.warnings.append(
                f"pagine senza testo {empty_pages} (probabile scansione): installare Tesseract "
                "e pytesseract per l'OCR, altrimenti il loro contenuto NON viene analizzato"
            )
    doc.text = "\n\n".join(texts)
    return doc


def _image(data: bytes) -> ExtractedDocument:
    if not ocr_available():
        raise UnsupportedDocument("le immagini richiedono l'OCR: installare Tesseract e pytesseract")
    from PIL import Image, ImageOps

    image = ImageOps.exif_transpose(Image.open(io.BytesIO(data)))  # foto da telefono ruotate
    return ExtractedDocument(_ocr(image), kind="immagine", pages=1,
                             warnings=["testo ottenuto con OCR: verificare, può sbagliare qualche carattere"])


def tesseract_cmd() -> str | None:
    found = shutil.which("tesseract")
    if found:
        return found
    return next((str(p) for p in _TESSERACT_CANDIDATES if p.is_file()), None)


def ocr_available() -> bool:
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return tesseract_cmd() is not None


def _ocr(image) -> str:
    import pytesseract

    pytesseract.pytesseract.tesseract_cmd = tesseract_cmd()
    langs = {"eng"}
    if TESSDATA_DIR.is_dir():
        # Variabile d'ambiente e non --tessdata-dir: su Windows pytesseract passerebbe
        # le virgolette del percorso letteralmente.
        os.environ["TESSDATA_PREFIX"] = str(TESSDATA_DIR)
        langs = {p.stem for p in TESSDATA_DIR.glob("*.traineddata")}
    lang = "ita+eng" if "ita" in langs else "eng"
    return pytesseract.image_to_string(image.convert("RGB"), lang=lang)


def _ocr_pages(data: bytes, pages: list[int]):
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    try:
        for i in pages:
            yield i, _ocr(pdf[i - 1].render(scale=300 / 72).to_pil())
    finally:
        pdf.close()


def _docx(data: bytes) -> ExtractedDocument:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    for section in document.sections:
        for part in (section.header, section.footer):
            parts.extend(p.text for p in part.paragraphs)
    return ExtractedDocument("\n".join(p for p in parts if p.strip()), kind="docx")
